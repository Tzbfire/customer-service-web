import asyncio
import io
import json
import time

import httpx
import pytest
from fastapi.testclient import TestClient
from PIL import Image

from desk.app import create_app
from desk.channels import decode_push, normalize_cookies
from desk.knowledge import Knowledge, split_chunks

H = {"X-Desk-Request": "1"}
PASSWORD = "a-test-password-123"


def completion(content):
    return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {"content": content}}]})


@pytest.fixture
def system(tmp_path):
    calls = []

    def model_handler(request):
        payload = json.loads(request.content)
        calls.append(payload)
        content = payload["messages"][-1]["content"]
        if isinstance(content, list):
            return completion("示例保温杯，容量 500mL，只能手洗，不可放入洗碗机。")
        if content == "请只回复 OK":
            return completion("OK")
        evidence = json.loads(content)["店铺资料"]
        return completion(
            json.dumps(
                {
                    "answer": "这款杯子容量为 500mL，需要手洗哦。",
                    "needs_human": False,
                    "source_ids": [evidence[0]["id"]],
                    "reason": "已找到对应的商品说明",
                },
                ensure_ascii=False,
            )
        )

    sent = []

    def channel_handler(request):
        if request.url.path == "/plateau/chat/list":
            return httpx.Response(200,json={"success":True,"result":{"messages":[],"has_more":False}})
        if request.url.path == "/latitude/user/userInfoWithTags":
            return httpx.Response(200, json={"success": False, "errorCode": 403})
        if request.url.path.endswith("userinfo"):
            return httpx.Response(200, json={"success": True, "result": {"id": "cs1", "mall_id": "mall1"}})
        if request.url.path.endswith("queryMerchantInfoByMallId"):
            return httpx.Response(
                200, json={"success": True, "result": {"mallId": "mall1", "mallName": "测试店铺"}}
            )
        if request.url.path.endswith("send_message"):
            sent.append(json.loads(request.content))
            return httpx.Response(200, json={"success": True, "result": {}})
        raise AssertionError("unexpected platform API path")

    app = create_app(tmp_path, httpx.MockTransport(model_handler), httpx.MockTransport(channel_handler))
    with TestClient(app) as client:
        result = client.post("/api/setup", headers=H, json={"username": "owner", "password": PASSWORD})
        assert result.status_code == 200
        app.state.store.save_model(
            {
                "base_url": "https://models.test/v1",
                "model": "text-test",
                "api_key": "test-only-secret",
                "vision_enabled": True,
                "vision_model": "vision-test",
            }
        )
        yield client, app, calls, sent


def upload(
    client, filename="cup.txt", content="杯子容量为 500mL，只能手洗，不可以放入洗碗机。", goods_id="100"
):
    raw = content.encode() if isinstance(content, str) else content
    result = client.post(
        "/api/documents",
        headers=H,
        data={"shop_id": 1, "goods_id": goods_id},
        files={"file": (filename, raw)},
    )
    assert result.status_code == 200, result.text
    document_id = result.json()["id"]
    for _ in range(100):
        value = client.get("/api/documents/" + document_id).json()
        if value["state"] != "parsing":
            return value
        time.sleep(0.01)
    raise AssertionError("parser timed out")


def publish(client, document):
    result = client.post(
        f"/api/documents/{document['id']}/publish",
        headers=H,
        json={
            "title": document["title"],
            "goods_id": document["goods_id"],
            "category": document["category"],
            "content": document["content"],
        },
    )
    assert result.status_code == 200, result.text


def receive(client, app, message_id="one", automatic=True):
    item = {
        "uid": "buyer1",
        "name": "买家一",
        "content": "杯子容量是多少？怎么清洗？",
        "goods_id": "100",
        "platform_id": message_id,
        "automatic": automatic,
    }
    client.portal.call(app.state.channels.receive, 1, item)
    return app.state.store.one("SELECT * FROM conversations WHERE buyer_uid=?", ("buyer1",))


def test_auth_setup_csrf_and_encrypted_key(system):
    client, app, _, _ = system
    assert (
        client.post("/api/setup", headers=H, json={"username": "again", "password": PASSWORD}).status_code
        == 409
    )
    assert client.post("/api/shops", json={"name": "bad"}).status_code == 403
    assert (
        client.post(
            "/api/shops", headers={**H, "Origin": "https://evil.test"}, json={"name": "bad"}
        ).status_code
        == 403
    )
    model = client.get("/api/model").json()
    assert model["has_key"] is True and "api_key" not in model
    stored = app.state.store.one("SELECT value FROM settings WHERE key='model'")["value"]
    assert "test-only-secret" not in stored
    client.post("/api/logout", headers=H)
    assert client.get("/api/documents?shop_id=1").status_code == 401
    assert client.get("/api/model").status_code == 401
    assert (
        client.post("/api/login", headers=H, json={"username": "owner", "password": PASSWORD}).status_code
        == 200
    )


def test_agent_can_answer_but_cannot_change_configuration(system):
    client, _, _, _ = system
    assert (
        client.post("/api/team", headers=H, json={"username": "agent1", "password": PASSWORD}).status_code
        == 200
    )
    client.post("/api/logout", headers=H)
    client.post("/api/login", headers=H, json={"username": "agent1", "password": PASSWORD})
    assert client.get("/api/shops").status_code == 200
    assert client.get("/api/model").status_code == 403
    assert (
        client.post("/api/team", headers=H, json={"username": "new", "password": PASSWORD}).status_code == 403
    )
    assert (
        client.post(
            "/api/documents", headers=H, data={"shop_id": 1}, files={"file": ("a.txt", b"a")}
        ).status_code
        == 403
    )
    answer = client.post("/api/playground", headers=H, json={"shop_id": 1, "question": "容量多少？"})
    assert answer.json()["needs_human"] is True


def test_disabled_member_session_is_revoked(system):
    client, app, _, _ = system
    member_id = client.post("/api/team", headers=H, json={"username": "agent1", "password": PASSWORD}).json()[
        "id"
    ]
    with TestClient(app) as second:
        second.post("/api/login", headers=H, json={"username": "agent1", "password": PASSWORD})
        assert second.get("/api/me").status_code == 200
        client.post(f"/api/team/{member_id}/disable", headers=H)
        assert second.get("/api/me").status_code == 401


def test_draft_is_not_searchable_until_published(system):
    client, app, calls, _ = system
    document = upload(client)
    assert document["state"] == "draft"
    answer = client.post(
        "/api/playground", headers=H, json={"shop_id": 1, "question": "杯子容量多少？", "goods_id": "100"}
    ).json()
    assert answer["sources"] == [] and answer["needs_human"] is True and not calls
    publish(client, document)
    answer = client.post(
        "/api/playground", headers=H, json={"shop_id": 1, "question": "杯子容量多少？", "goods_id": "100"}
    ).json()
    assert answer["sources"][0]["document_id"] == document["id"]
    assert answer["needs_human"] is False and "500mL" in answer["answer"]
    assert calls[-1]["model"] == "text-test"
    assert app.state.store.one("SELECT COUNT(*) n FROM chunks")["n"] > 0


def test_shop_and_product_scopes_are_respected(system):
    client, app, _, _ = system
    publish(client, upload(client))
    other_id = client.post("/api/shops", headers=H, json={"name": "另一个店铺"}).json()["id"]
    assert not app.state.knowledge.search(other_id, "杯子容量", "100")
    assert not app.state.knowledge.search(1, "杯子容量", "200")
    policy = upload(client, "policy.txt", "本店商品支持七天退换货，退货前请先联系客服。", "")
    publish(client, policy)
    result = app.state.knowledge.search(1, "退换货", "200")
    assert len(result) == 1 and result[0]["goods_id"] == ""


def test_image_ocr_retains_review_step(system):
    client, _, calls, _ = system
    image = io.BytesIO()
    Image.new("RGB", (40, 40), "white").save(image, "PNG")
    document = upload(client, "spec.png", image.getvalue())
    assert document["state"] == "draft" and "500mL" in document["content"]
    assert calls[-1]["model"] == "vision-test"
    assert calls[-1]["messages"][-1]["content"][1]["image_url"]["url"].startswith("data:image/jpeg;base64,")


def test_image_without_vision_fails_with_actionable_message(system):
    client, app, calls, _ = system
    model = app.state.store.model()
    model["vision_enabled"] = False
    app.state.store.save_model(model)
    image = io.BytesIO()
    Image.new("RGB", (40, 40), "white").save(image, "PNG")
    document = upload(client, "spec.png", image.getvalue())
    assert document["state"] == "failed"
    assert "图片识别" in document["error"] and not calls


def test_duplicate_and_unsupported_uploads(system):
    client, _, _, _ = system
    upload(client)
    result = client.post(
        "/api/documents",
        headers=H,
        data={"shop_id": 1, "goods_id": "100"},
        files={"file": ("same.txt", "杯子容量为 500mL，只能手洗，不可以放入洗碗机。".encode())},
    )
    assert result.status_code == 409
    result = client.post(
        "/api/documents", headers=H, data={"shop_id": 1}, files={"file": ("a.html", b"<script/>")}
    )
    assert result.status_code == 400


def test_deleting_document_removes_search_chunks(system):
    client, app, _, _ = system
    document = upload(client)
    publish(client, document)
    assert client.delete("/api/documents/" + document["id"], headers=H).status_code == 200
    assert not app.state.knowledge.search(1, "容量")
    assert not (app.state.store.directory / "uploads" / document["id"]).exists()


def test_word_excel_and_pdf_extraction(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    import pymupdf

    doc = Document()
    doc.add_paragraph("保温杯")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "容量"
    table.cell(0, 1).text = "500mL"
    doc.add_paragraph("只能手洗")
    path = tmp_path / "spec.docx"
    doc.save(path)
    text = Knowledge.extract_offline(path, ".docx")
    assert text.index("保温杯") < text.index("500mL") < text.index("只能手洗")
    book = Workbook()
    book.active.append(["商品ID", "容量"])
    book.active.append(["100", "500mL"])
    path = tmp_path / "spec.xlsx"
    book.save(path)
    assert "100 | 500mL" in Knowledge.extract_offline(path, ".xlsx")
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((40, 60), "Cup product: Capacity 500mL. Handwash only. Do not use dishwasher.")
    path = tmp_path / "spec.pdf"
    pdf.save(path)
    from desk.store import Store

    knowledge = Knowledge(Store(tmp_path / "data"), None)
    assert "500mL" in asyncio.run(knowledge.extract(path, ".pdf"))


def test_review_and_dedup_do_not_send_messages(system):
    client, app, _, sent = system
    publish(client, upload(client))
    conversation = receive(client, app)
    receive(client, app)
    messages = app.state.store.all("SELECT * FROM messages WHERE conversation_id=?", (conversation["id"],))
    assert len(messages) == 2 and [item["role"] for item in messages] == ["customer", "draft"]
    assert conversation["needs_human"] == 1 and not sent


def test_auto_send_requires_sources_and_platform_ack(system):
    client, app, _, sent = system
    publish(client, upload(client))
    app.state.store.execute(
        "UPDATE shops SET cookies=?,mode=? WHERE id=1",
        (app.state.store.encrypt(json.dumps({"test": "cookie"})), "auto"),
    )
    conversation = receive(client, app)
    assert len(sent) == 1 and sent[0]["data"]["message"]["to"]["uid"] == "buyer1"
    assert conversation["needs_human"] == 0
    receive(client, app)
    assert len(sent) == 1
    receive(client, app, "two", automatic=False)
    assert len(sent) == 1


def test_human_takeover_prevents_ai_and_rejects_other_agent(system):
    client, app, _, sent = system
    conversation = receive(client, app)
    identifier = conversation["id"]
    assert client.post(f"/api/conversations/{identifier}/takeover", headers=H).status_code == 200
    receive(client, app, "two")
    assert app.state.store.one("SELECT COUNT(*) n FROM messages")["n"] == 3
    client.post("/api/team", headers=H, json={"username": "agent1", "password": PASSWORD})
    with TestClient(app) as second:
        second.post("/api/login", headers=H, json={"username": "agent1", "password": PASSWORD})
        assert second.post(f"/api/conversations/{identifier}/takeover", headers=H).status_code == 409
        assert (
            second.post(
                f"/api/conversations/{identifier}/send",
                headers=H,
                json={"content": "bad", "request_key": "unique-123"},
            ).status_code
            == 409
        )
    assert not sent


def test_manual_send_is_idempotent(system):
    client, app, _, sent = system
    identifier = receive(client, app)["id"]
    app.state.store.execute(
        "UPDATE shops SET cookies=? WHERE id=1", (app.state.store.encrypt(json.dumps({"test": "cookie"})),)
    )
    client.post(f"/api/conversations/{identifier}/takeover", headers=H)
    payload = {"content": "您好，请问有什么可以帮您？", "request_key": "same-request-123"}
    first = client.post(f"/api/conversations/{identifier}/send", headers=H, json=payload)
    second = client.post(f"/api/conversations/{identifier}/send", headers=H, json=payload)
    assert first.status_code == second.status_code == 200
    assert first.json()["id"] == second.json()["id"] and len(sent) == 1


def test_decode_only_buyer_push_and_keep_goods_scope():
    base = {
        "response": "push",
        "message": {
            "type": 0,
            "content": "怎么清洗？",
            "msg_id": "1",
            "from": {"role": "user", "uid": "buyer"},
            "info": {"goodsID": 100, "goodsName": "杯子"},
        },
    }
    item = decode_push(base)
    assert item["goods_id"] == "100" and item["automatic"] is True
    base["message"]["from"]["role"] = "mall_cs"
    assert decode_push(base) is None
    base["message"]["from"]["role"] = "user"
    base["message"]["type"] = 1
    assert decode_push(base)["automatic"] is False
    assert decode_push({"response": "auth", "auth": {"result": "ok"}}) is None
    with pytest.raises(ValueError):
        normalize_cookies("{bad json}")


def test_chunks_cover_tail_and_do_not_loop():
    text = "参数为500mL。\n" * 500 + "END"
    chunks = split_chunks(text)
    assert chunks[-1].endswith("END") and all(len(part) <= 900 for part in chunks)


def test_fabricated_citation_cannot_enable_auto_send(system):
    client, app, _, _ = system
    publish(client, upload(client))

    async def bad_response(*args, **kwargs):
        return json.dumps({"answer": "随便编造", "needs_human": False, "source_ids": [999999]})

    app.state.knowledge.model.complete = bad_response
    answer = client.post(
        "/api/playground", headers=H, json={"shop_id": 1, "question": "杯子容量", "goods_id": "100"}
    ).json()
    assert answer["needs_human"] is True and answer["sources"] == []


def test_long_image_keeps_readable_tiles():
    from desk.knowledge import image_tiles
    import base64

    data = io.BytesIO()
    Image.new("RGB", (700, 6500), "white").save(data, "PNG")
    tiles = image_tiles(data.getvalue())
    assert len(tiles) == 4
    with Image.open(io.BytesIO(base64.b64decode(tiles[0]))) as tile:
        assert tile.width == 700 and tile.height == 1800


def test_word_embedded_image_uses_vision(system, tmp_path):
    from docx import Document

    client, _, calls, _ = system
    image = tmp_path / "photo.png"
    Image.new("RGB", (100, 100), "white").save(image)
    document = Document()
    document.add_paragraph("测试商品说明")
    document.add_picture(str(image))
    path = tmp_path / "mixed.docx"
    document.save(path)
    result = upload(client, "mixed.docx", path.read_bytes())
    assert result["state"] == "draft" and "内嵌图片" in result["content"]
    assert "500mL" in result["content"] and calls[-1]["model"] == "vision-test"


def test_multiple_products_require_clarification(system):
    client, _, calls, _ = system
    publish(client, upload(client))
    publish(client, upload(client, "other.txt", "另一个杯子容量为 1000mL，只能手洗。", "200"))
    result = client.post(
        "/api/playground", headers=H, json={"shop_id": 1, "question": "杯子容量是多少？"}
    ).json()
    assert result["needs_human"] is True and "哪款商品" in result["answer"]
    assert not calls


def test_takeover_during_generation_does_not_block_manual_reply(system):
    client, app, _, sent = system
    app.state.store.execute(
        "UPDATE shops SET cookies=?,mode=? WHERE id=1",
        (app.state.store.encrypt(json.dumps({"test": "cookie"})), "auto"),
    )

    async def scenario():
        started, release = asyncio.Event(), asyncio.Event()

        async def slow_answer(*args, **kwargs):
            started.set()
            await release.wait()
            return {"answer": "stale AI response", "needs_human": False, "sources": [{"id": 1}]}

        app.state.knowledge.answer = slow_answer
        item = {
            "uid": "buyer1",
            "name": "buyer",
            "goods_id": "100",
            "content": "容量多少",
            "platform_id": "slow-one",
            "automatic": True,
        }
        task = asyncio.create_task(app.state.channels.receive(1, item))
        await asyncio.wait_for(started.wait(), 2)
        conversation = app.state.store.one("SELECT * FROM conversations WHERE buyer_uid='buyer1'")
        app.state.store.execute(
            "UPDATE conversations SET mode='human',assignee=1 WHERE id=?", (conversation["id"],)
        )

        async def manual_reply():
            async with app.state.channels.lock(conversation["id"]):
                await app.state.channels.send(conversation["id"], "人工马上为您核实", 1, "takeover-manual")

        await asyncio.wait_for(manual_reply(), 2)
        release.set()
        await task

    client.portal.call(scenario)
    assert len(sent) == 1 and sent[0]["data"]["message"]["content"] == "人工马上为您核实"
    assert not app.state.store.one("SELECT id FROM messages WHERE role='draft'")


def test_network_failure_does_not_retry_sending(system):
    client, app, _, _ = system
    identifier = receive(client, app)["id"]
    app.state.store.execute(
        "UPDATE shops SET cookies=? WHERE id=1", (app.state.store.encrypt(json.dumps({"test": "cookie"})),)
    )
    client.post(f"/api/conversations/{identifier}/takeover", headers=H)
    requests = []

    def fail(request):
        requests.append(request)
        raise httpx.ReadTimeout("test timeout")

    app.state.channels.transport = httpx.MockTransport(fail)
    payload = {"content": "test", "request_key": "uncertain-123"}
    assert client.post(f"/api/conversations/{identifier}/send", headers=H, json=payload).status_code == 502
    assert client.post(f"/api/conversations/{identifier}/send", headers=H, json=payload).status_code == 409
    assert len(requests) == 1
    assert app.state.store.one("SELECT state FROM messages WHERE role='agent'")["state"] == "uncertain"


def test_model_error_cannot_be_used_as_a_reply_draft(system):
    from desk.llm import ModelError

    client, app, _, sent = system

    async def fail(*args, **kwargs):
        raise ModelError("模型暂时不可用")

    app.state.knowledge.answer = fail
    conversation = receive(client, app)
    assert conversation["needs_human"] == 1
    assert app.state.store.one("SELECT id FROM messages WHERE role='system'")
    assert not app.state.store.one("SELECT id FROM messages WHERE role='draft'")
    assert not sent


def test_binding_uses_verified_platform_identity(system):
    client, app, _, _ = system
    response = client.post("/api/shops/1/bind", headers=H, json={"cookies": json.dumps({"test": "cookie"})})
    assert response.status_code == 200
    shop = client.get("/api/shops").json()[0]
    assert shop["platform_shop_id"] == "mall1" and shop["platform_user_id"] == "cs1"
    assert "cookies" not in shop
    saved = app.state.store.one("SELECT cookies FROM shops WHERE id=1")["cookies"]
    assert "cookie" not in saved


def test_model_endpoint_requires_explicit_private_opt_in(system):
    client, app, _, _ = system
    value = {"base_url": "http://127.0.0.1:12345/v1", "model": "local-test"}
    assert client.post("/api/model", headers=H, json=value).status_code == 400
    value["allow_private"] = True
    assert client.post("/api/model", headers=H, json=value).status_code == 200
    assert app.state.store.model()["api_key"] == "test-only-secret"
    value["base_url"] = "http://169.254.169.254/"
    assert client.post("/api/model", headers=H, json=value).status_code == 400


def test_closing_login_cancels_the_owned_browser_task(system):
    client, app, _, _ = system
    stopped = []

    async def prepare():
        async def wait_for_user():
            try:
                await asyncio.Event().wait()
            finally:
                stopped.append(True)

        app.state.channels.login_tasks[1] = asyncio.create_task(wait_for_user())
        await asyncio.sleep(0)

    client.portal.call(prepare)
    assert client.post("/api/shops/1/login-cancel", headers=H).status_code == 200
    assert stopped == [True]
    assert 1 not in app.state.channels.login_tasks
    assert client.get("/api/shops").json()[0]["login"]["status"] == "cancelled"


def test_login_uses_a_visible_normal_browser_and_saves_only_after_binding(system, monkeypatch):
    from types import SimpleNamespace
    import playwright.async_api

    client, app, _, _ = system
    options = []
    calls = []

    class Page:
        url = "about:blank"

        async def goto(self, url, **kwargs):
            self.url = url

        async def bring_to_front(self):
            calls.append("foreground")

        def is_closed(self):
            return False

    page = Page()

    class Context:
        pages = [page]

        async def cookies(self):
            return [{"name": "fixture", "value": "test", "domain": ".pinduoduo.com"}]

        async def close(self):
            calls.append("closed")

    class Chromium:
        async def launch_persistent_context(self, *args, **kwargs):
            options.append(kwargs)
            return Context()

    class Runtime:
        async def __aenter__(self):
            return SimpleNamespace(chromium=Chromium())

        async def __aexit__(self, *args):
            return False

    monkeypatch.setattr(playwright.async_api, "async_playwright", lambda: Runtime())

    async def bind(shop_id, cookies):
        calls.append("bound")

    app.state.channels.bind = bind
    client.portal.call(app.state.channels._login, 1)
    assert options[0]["headless"] is False
    assert options[0]["no_viewport"] is True
    assert calls.index("foreground") < calls.index("bound") < calls.index("closed")
    assert app.state.channels.login_states[1]["status"] == "done"


def test_native_login_can_focus_an_existing_window(system):

    client, app, _, _ = system
    calls = []

    class Page:
        def is_closed(self):
            return False

        async def bring_to_front(self):
            calls.append("foreground")

    app.state.channels.login_pages[1] = Page()
    assert client.post("/api/shops/1/login-focus", headers=H).status_code == 200
    assert calls == ["foreground"]
    app.state.channels.login_pages.clear()
    assert client.post("/api/shops/1/login-focus", headers=H).status_code == 409


def test_login_redirect_query_is_not_mistaken_for_an_authenticated_workstation(system, monkeypatch):
    from types import SimpleNamespace
    import playwright.async_api

    client, app, _, _ = system
    calls = []

    class Page:
        url = "about:blank"
        checks = 0

        async def goto(self, url, **kwargs):
            self.url = "https://mms.pinduoduo.com/login/?redirectUrl=https%3A%2F%2Fmms.pinduoduo.com%2Fchat-merchant%2Findex.html"

        async def bring_to_front(self):
            pass

        def is_closed(self):
            self.checks += 1
            return self.checks > 1

    page = Page()

    class Context:
        pages = [page]

        async def cookies(self):
            calls.append("read_cookies")
            return [{"name": "fixture", "value": "pre-login", "domain": ".pinduoduo.com"}]

        async def close(self):
            pass

    class Chromium:
        async def launch_persistent_context(self, *args, **kwargs):
            return Context()

    class Runtime:
        async def __aenter__(self):
            return SimpleNamespace(chromium=Chromium())

        async def __aexit__(self, *args):
            return False

    monkeypatch.setattr(playwright.async_api, "async_playwright", lambda: Runtime())

    async def bind(*args):
        calls.append("bind")

    app.state.channels.bind = bind
    client.portal.call(app.state.channels._login, 1)
    assert calls == [], "Login redirects must not trigger a premature identity read."
    assert app.state.channels.login_states[1]["status"] == "cancelled"
